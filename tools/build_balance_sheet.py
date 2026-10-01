"""Builds doc/Road_to_Jerusalem_Balance_Sheet.xlsx - the game design balance
sheet - from the game's own data (characters, classes, skills) and the
balance constants in scripts/game_logic.py.

Every derived number in the workbook is a live Excel formula that points at
the Constants sheet, so changing a constant there shows its effect across
every unit, class and matchup. The workbook is a sandbox: the game never
reads it - copy a change you like back into game_logic.py or the CSVs.

Run from the project root:   python tools/build_balance_sheet.py
"""
import csv
import glob
import os
import sys

os.environ.setdefault("SDL_VIDEODRIVER", "dummy")
os.environ.setdefault("SDL_AUDIODRIVER", "dummy")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from openpyxl import Workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

import scripts.game_logic as g  # noqa: E402
from scripts.config import FAITH_CAP  # noqa: E402
from scripts.data_editor import load_characters_from_csv, load_skills_from_csv, load_stage_manifest  # noqa: E402
from scripts.hero import load_classes_from_csv  # noqa: E402

OUT = os.path.join(ROOT, "doc", "Road_to_Jerusalem_Balance_Sheet.xlsx")

FONT = "Arial"
INPUT = Font(name=FONT, color="0000FF")          # values copied from the game
CALC = Font(name=FONT, color="000000")           # formulas
BOLD = Font(name=FONT, bold=True)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, bold=True, size=16)
HEAD_FILL = PatternFill("solid", fgColor="5B3A1E")
SECTION_FILL = PatternFill("solid", fgColor="EADBC8")
EDIT_FILL = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="C8B8A0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PCT = "0.0%"
NUM1 = "0.0"


def header(ws, row, labels, widths=None):
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font, cell.fill, cell.border = HEAD, HEAD_FILL, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 62
    for col, width in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def put(ws, row, col, value, font=CALC, fmt=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font, cell.border = font, BOX
    if fmt:
        cell.number_format = fmt
    return cell


# --- Constants -------------------------------------------------------------------

def constants_rows():
    """(section, code name, value, format, what it does, source)."""
    gl = "scripts/game_logic.py"
    c = []

    def add(section, name, value, fmt, what, source=None):
        c.append((section, name, value, fmt, what, source or f"{gl}: {name}"))

    add("Faith", "FAITH_CAP", FAITH_CAP, "0", "Faith ceiling. An enemy preached up to it converts.", "scripts/config.py: FAITH_CAP")
    add("Faith", "FAITH_TRANSFER_RATE", g.FAITH_TRANSFER_RATE, "0.00", "Preach/Heal strength = (preacher Faith + Speech) x this rate.")
    add("Faith", "MISSIONARY_PREACH_MULTIPLIER", g.CLASS_PREACH_MULTIPLIER["Missionary"], "0.00", "Missionary (Paul) preaches this many times stronger.", f"{gl}: CLASS_PREACH_MULTIPLIER['Missionary']")
    add("Faith", "FAITH_BASE_HIT_CHANCE", g.FAITH_BASE_HIT_CHANCE, PCT, "Preach hit chance before Love.")
    add("Faith", "FAITH_HIT_PER_LOVE", g.FAITH_HIT_PER_LOVE, "0.000", "Preach hit chance added per point of the preacher's Love.")
    add("Faith", "FAITH_MIN_HIT_CHANCE", g.FAITH_MIN_HIT_CHANCE, PCT, "Preach hit chance floor.")
    add("Faith", "FAITH_MAX_HIT_CHANCE", g.FAITH_MAX_HIT_CHANCE, PCT, "Preach hit chance ceiling.")
    add("Faith", "FAITH_DECAY_PER_TURN", 0.01, PCT, "Every unit loses this share of its faith at the start of its own turn.", f"{gl}: begin_turn (hard-coded 0.01)")
    add("Faith", "SIDE_SWITCH_FAITH", g.SIDE_SWITCH_FAITH, "0", "Faith a unit has right after converting or turning back.")
    add("Despair", "DESPAIR_TURNS", g.DESPAIR_TURNS, "0", "Turns a party member at 0 faith sits out.")
    add("Despair", "DESPAIR_RECOVERY_FAITH", g.DESPAIR_RECOVERY_FAITH, "0", "Faith they recover with when Despair lifts.")
    add("Physical", "PHYSICAL_BASE_HIT_CHANCE", g.PHYSICAL_BASE_HIT_CHANCE, PCT, "Slash/Shoot/Sling hit chance on level ground before modifiers.")
    add("Physical", "ELEVATION_HIT_BONUS_PER_TILE", g.ELEVATION_HIT_BONUS_PER_TILE, PCT, "Per tile of height the attacker stands above (+) or below (-) the target. Archers ignore the uphill penalty.")
    add("Physical", "PHYSICAL_MIN_HIT_CHANCE", g.PHYSICAL_MIN_HIT_CHANCE, PCT, "Physical hit chance floor.")
    add("Physical", "PHYSICAL_MAX_HIT_CHANCE", g.PHYSICAL_MAX_HIT_CHANCE, PCT, "Physical hit chance ceiling.")
    add("Physical", "MAGIC_DEFENSE_HIT_SCALE", g.MAGIC_DEFENSE_HIT_SCALE, "0.000", "Hit chance removed per point of the target's Resist.")
    add("Physical", "SOLDIER_HIT_BONUS", g.CLASS_HIT_BONUS["Soldier"], PCT, "Soldier and Sergeant attackers hit this much more often.", f"{gl}: CLASS_HIT_BONUS")
    add("Physical", "SHIELDBEARER_HIT_PENALTY", g.CLASS_HIT_PENALTY_AGAINST["Shieldbearer"], PCT, "Physical hits against a Shieldbearer are this much less likely.", f"{gl}: CLASS_HIT_PENALTY_AGAINST")
    add("Physical", "OFFICER_AURA_HIT_BONUS", g.OFFICER_AURA_HIT_BONUS, PCT, "Hit bonus for allies near an Officer (Centurion Marcus).")
    add("Physical", "OFFICER_AURA_RANGE", g.OFFICER_AURA_RANGE, "0", "Tiles the Officer's aura reaches.")
    add("Physical", "DISARM_TURNS", g.DISARM_TURNS, "0", "Turns a Slash or Sling hit stops the target from acting.")
    add("Status", "PATIENCE_RESIST_SCALE", g.PATIENCE_RESIST_SCALE, "0.000", "Chance per point of Patience to shrug off Snare/Stun/Disarm skills.")
    add("Status", "PATIENCE_RESIST_CAP", g.PATIENCE_RESIST_CAP, PCT, "Status resist ceiling.")
    add("Status", "SHOVE_CHANCE", g.SHOVE_CHANCE, PCT, "Shove's base chance to stun and push (before resist).")
    add("Status", "SHOVE_DISTANCE", g.SHOVE_DISTANCE, "0", "Tiles a Shove pushes (Sergeants: 3).")
    add("Status", "FISH_NET_SNARE_TURNS", 2, "0", "Turns a Fish net snares its target.", f"{gl}: _resolve_skill_status (hard-coded 2)")
    add("Status", "COMMAND_CT_BOOST", g.COMMAND_CT_BOOST, "0", "CT Command adds to an ally (100 CT = a turn).")
    add("Support", "LEAD_LOVE_BONUS", g.LEAD_LOVE_BONUS, "0", "Love Lead gives every ally within range.")
    add("Support", "LEAD_RADIUS", g.LEAD_RADIUS, "0", "Tiles Lead reaches.")
    add("Support", "LEAD_TURNS", g.LEAD_TURNS, "0", "Turns Lead's Love lasts.")
    add("Support", "RALLY_FAITH_RESTORE", g.RALLY_FAITH_RESTORE, "0", "Faith Rally moves (enemy: down toward the baseline; party: up).")
    add("Support", "RALLY_BASELINE_FAITH", g.RALLY_BASELINE_FAITH, "0", "Enemy Rally never lowers faith below this.")
    add("Morale", "MORALE_BASE_CHANCE", g.MORALE_BASE_CHANCE, PCT, "Chance of a morale surge each turn, before Bravery (once per battle).")
    add("Morale", "MORALE_CHANCE_SCALE", g.MORALE_CHANCE_SCALE, "0.000", "Surge chance per point of Bravery.")
    add("Morale", "MORALE_CHANCE_CAP", g.MORALE_CHANCE_CAP, PCT, "Surge chance ceiling.")
    add("Morale", "MORALE_FAITH_BONUS", g.MORALE_FAITH_BONUS, "0", "Faith a surge adds (for that turn).")
    add("Morale", "MORALE_MV_BONUS", g.MORALE_MV_BONUS, "0", "Move a surge adds (for that turn).")
    add("Morale", "MORALE_SPEED_BONUS", g.MORALE_SPEED_BONUS, "0", "Speed a surge adds (for that turn).")
    add("Fleeing", "FLEE_MAX_CHANCE", g.FLEE_MAX_CHANCE, PCT, "Legionnaire flee chance per turn once their whole force is gone, before courage.")
    add("Fleeing", "FLEE_COURAGE_SCALE", g.FLEE_COURAGE_SCALE, "0", "Bravery + Patience that makes a legionnaire never flee.")
    add("Persecute", "PERSECUTE_FAITH", g.PERSECUTE_FAITH, "0", "Fixed faith Persecute takes (Damascus). At 0 the disciple is arrested.")
    add("Persecute", "PERSECUTE_SLOW", g.PERSECUTE_SLOW, "0", "Move Persecute takes away.")
    add("Persecute", "PERSECUTE_TURNS", g.PERSECUTE_TURNS, "0", "Turns the slow lasts (refreshed, not stacked).")
    add("Mounts", "HORSE_MOVE_BONUS", g.MOUNTS["horse"]["mv"], "0", "Move a horse adds.", f"{gl}: MOUNTS['horse']['mv']")
    add("Turn order", "CT_PER_TURN", 100, "0", "CT needed to act. Each tick a unit gains CT equal to its Speed.", f"{gl}: battle loop (hard-coded 100)")
    add("EXP", "EXP_BASE", g.EXP_BASE, "0", "EXP for an action on a same-level target (+/- level gap).")
    add("EXP", "EXP_MIN", g.EXP_MIN, "0", "Action EXP floor.")
    add("EXP", "EXP_MAX", g.EXP_MAX, "0", "Action EXP ceiling.")
    add("EXP", "CONVERT_EXP_BASE", g.CONVERT_EXP_BASE, "0", "EXP for the action that converts an enemy (+2 per level gap).")
    add("EXP", "EXP_PER_LEVEL", g.EXP_PER_LEVEL, "0", "EXP needed per level.")
    add("EXP", "SPEED_GROWTH_EVERY", g.SPEED_GROWTH_EVERY, "0", "A unit gains +1 Speed on every Nth level.")
    for name, value in [("REASON_TOGETHER_FAITH", g.REASON_TOGETHER_FAITH), ("INTERCESSION_FAITH", g.INTERCESSION_FAITH),
                        ("INSPIRE_FAITH", g.INSPIRE_FAITH), ("GOOD_NEWS_SHARE", g.GOOD_NEWS_SHARE), ("MARSHAL_CT", g.MARSHAL_CT),
                        ("PSALM_CT", g.PSALM_CT), ("BOLD_VENTURE_CHANCE", g.BOLD_VENTURE_CHANCE), ("BOLD_VENTURE_FAITH", g.BOLD_VENTURE_FAITH),
                        ("BOLD_VENTURE_BACKFIRE", g.BOLD_VENTURE_BACKFIRE), ("PARABLE_FAITH", g.PARABLE_FAITH), ("PARABLE_CT", g.PARABLE_CT),
                        ("BREAKING_BREAD_FAITH", g.BREAKING_BREAD_FAITH), ("ARREST_SNARE_TURNS", g.ARREST_SNARE_TURNS)]:
        add("Class skills", name, value, PCT if isinstance(value, float) and value < 1 else "0", "Signature-skill number (see Skills sheet).")
    return c


def build_constants(wb):
    ws = wb.create_sheet("Constants")
    header(ws, 1, ["Section", "Constant", "Value", "What it does", "Source in the game"], [14, 32, 11, 70, 52])
    ref = {}
    for row, (section, name, value, fmt, what, source) in enumerate(constants_rows(), start=2):
        put(ws, row, 1, section)
        put(ws, row, 2, name, BOLD)
        put(ws, row, 3, value, INPUT, fmt)
        put(ws, row, 4, what)
        put(ws, row, 5, source)
        ref[name] = f"Constants!$C${row}"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:E{ws.max_row}"
    return ref


# --- Level growth ----------------------------------------------------------------

def build_growth(wb):
    ws = wb.create_sheet("Level Growth")
    put(ws, 1, 1, "Level to preview", BOLD)
    level = put(ws, 1, 2, 10, INPUT, "0")
    level.fill = EDIT_FILL
    level.comment = Comment("Type any level. The 'gained by level' columns show what a Lv 1 unit of each class has gained by then.", "Balance sheet")
    header(ws, 3, ["Class", "Max MP / level", "Speech / level", "Resist / level", "MP gained by level", "Speech gained by level", "Resist gained by level", "Speed gained by level"],
           [16, 14, 14, 14, 16, 18, 18, 18])
    classes = list(g.LEVEL_GROWTH.items()) + [("(any other class)", g.DEFAULT_LEVEL_GROWTH)]
    speed_every = "Constants!$C$" + str(_constant_row("SPEED_GROWTH_EVERY"))
    for row, (name, growth) in enumerate(classes, start=4):
        put(ws, row, 1, name, BOLD)
        put(ws, row, 2, growth["max_mp"], INPUT)
        put(ws, row, 3, growth["magic_attack"], INPUT)
        put(ws, row, 4, growth["magic_defense"], INPUT)
        for col, src in ((5, "B"), (6, "C"), (7, "D")):
            put(ws, row, col, f"={src}{row}*($B$1-1)")
        put(ws, row, 8, f"=INT($B$1/{speed_every})")
    put(ws, row + 2, 1, f"Source: scripts/game_logic.py LEVEL_GROWTH / DEFAULT_LEVEL_GROWTH. Speed: +1 on every level divisible by SPEED_GROWTH_EVERY.")
    ws.freeze_panes = "B4"
    return ws, 4, row  # first/last growth rows


_CONSTANT_ROWS = {}


def _constant_row(name):
    return _CONSTANT_ROWS[name]


# --- Units -----------------------------------------------------------------------

UNIT_COLUMNS = ["Key (stage - name)", "Stage", "Name", "Team", "Class", "Level", "Mount", "Skills",
                "Speed", "Move", "Jump", "MP", "Speech (magic_attack)", "Resist (magic_defense)", "Faith", "Bravery", "Patience", "Love",
                "Eff. Speech at level", "Eff. Resist at level", "Eff. Speed at level", "Eff. Move (incl. mount)",
                "Preach strength (faith per hit)", "Preach hit %", "Expected faith / preach", "Physical hit % against them (flat)",
                "Status resist %", "Morale surge % / turn", "Ticks per turn", "Flee % / turn at half force lost", "Persecutes to arrest"]


def growth_lookup(growth_ws, first, last, class_cell, column):
    """Per-level growth for a class, falling back to the default row."""
    rng = f"'Level Growth'!$A${first}:$A${last}"
    col = f"'Level Growth'!${column}${first}:${column}${last}"
    default = f"'Level Growth'!${column}${last}"
    return f"IFERROR(INDEX({col},MATCH({class_cell},{rng},0)),{default})"


def build_units(wb, ref, growth):
    growth_ws, gfirst, glast = growth
    ws = wb.create_sheet("Units")
    header(ws, 1, UNIT_COLUMNS, [26, 11, 24, 8, 13, 7, 7, 30] + [8] * 10 + [10, 10, 10, 11, 13, 10, 13, 13, 11, 11, 9, 13, 11])
    manifest = load_stage_manifest()
    rows = []
    for node in manifest:
        for c in load_characters_from_csv(manifest[node]["characters"]):
            rows.append((node, c))
    C = ref
    for row, (stage, c) in enumerate(rows, start=2):
        put(ws, row, 1, f'=B{row}&" - "&C{row}')
        for col, value in enumerate([stage, c["name"], c["team"], c["class"], c["level"], c.get("mount", ""), "|".join(c["skills"]),
                                     c["speed"], c["mv"], c["jump"], c["mp"], c["magic_attack"], c["magic_defense"],
                                     c["faith"], c["bravery"], c["patience"], c["love"]], start=2):
            put(ws, row, col, value, INPUT)
        grow = lambda column: growth_lookup(growth_ws, gfirst, glast, f"E{row}", column)  # noqa: E731
        put(ws, row, 19, f"=M{row}+{grow('C')}*(F{row}-1)")
        put(ws, row, 20, f"=N{row}+{grow('D')}*(F{row}-1)")
        put(ws, row, 21, f"=I{row}+INT(F{row}/{C['SPEED_GROWTH_EVERY']})")
        put(ws, row, 22, f'=J{row}+IF(G{row}="horse",{C["HORSE_MOVE_BONUS"]},0)')
        put(ws, row, 23, f'=(O{row}+S{row})*{C["FAITH_TRANSFER_RATE"]}*IF(E{row}="Missionary",{C["MISSIONARY_PREACH_MULTIPLIER"]},1)', fmt=NUM1)
        put(ws, row, 24, f"=MIN({C['FAITH_MAX_HIT_CHANCE']},MAX({C['FAITH_MIN_HIT_CHANCE']},{C['FAITH_BASE_HIT_CHANCE']}+R{row}*{C['FAITH_HIT_PER_LOVE']}))", fmt=PCT)
        put(ws, row, 25, f"=W{row}*X{row}", fmt=NUM1)
        put(ws, row, 26, f'=MIN({C["PHYSICAL_MAX_HIT_CHANCE"]},MAX({C["PHYSICAL_MIN_HIT_CHANCE"]},{C["PHYSICAL_BASE_HIT_CHANCE"]}-T{row}*{C["MAGIC_DEFENSE_HIT_SCALE"]}-IF(E{row}="Shieldbearer",{C["SHIELDBEARER_HIT_PENALTY"]},0)))', fmt=PCT)
        put(ws, row, 27, f"=MIN({C['PATIENCE_RESIST_CAP']},Q{row}*{C['PATIENCE_RESIST_SCALE']})", fmt=PCT)
        put(ws, row, 28, f"=MIN({C['MORALE_CHANCE_CAP']},{C['MORALE_BASE_CHANCE']}+P{row}*{C['MORALE_CHANCE_SCALE']})", fmt=PCT)
        put(ws, row, 29, f"={C['CT_PER_TURN']}/U{row}", fmt=NUM1)
        put(ws, row, 30, f'=IF(LEFT(C{row},11)="Legionnaire",{C["FLEE_MAX_CHANCE"]}*0.5*MAX(0,1-(P{row}+Q{row})/{C["FLEE_COURAGE_SCALE"]}),"-")', fmt=PCT)
        put(ws, row, 31, f'=IF(D{row}="Enemy",ROUNDUP(O{row}/{C["PERSECUTE_FAITH"]},0),"-")')
    last = len(rows) + 1
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(UNIT_COLUMNS))}{last}"
    notes = {
        "W1": "(Faith + Eff. Speech) x FAITH_TRANSFER_RATE, x1.5 for the Missionary. Party preaching an enemy adds this; enemies preaching the party take it away.",
        "X1": "FAITH_BASE_HIT_CHANCE + Love x FAITH_HIT_PER_LOVE, kept within the min/max. Heals on allies always land.",
        "Z1": "Chance a Slash/Shoot/Sling lands on this unit on level ground from an attacker with no class bonus: base - Eff. Resist x scale, -15% for Shieldbearers.",
        "AA1": "Chance this unit shrugs off Fish net, Shove's stun, Arrest, Peacemaker, Disputation.",
        "AC1": "Ticks of the CT clock between this unit's turns (lower = acts more often).",
        "AD1": "Legionnaires only: chance per turn once half their original force has converted or fled.",
        "AE1": "Enemies only: how many Persecutes take their faith to 0 (arrest) from where they start.",
        "S1": "Stats at the unit's level: CSV base + Level Growth x (level - 1). Equipment and books are not included.",
    }
    for cell, text in notes.items():
        ws[cell].comment = Comment(text, "Balance sheet")
    return ws, last


# --- Classes ---------------------------------------------------------------------

def build_classes(wb, ref):
    ws = wb.create_sheet("Classes")
    labels = ["Type", "Class", "Role", "Signature skill", "Skills", "Speed", "Move", "Jump", "MP", "Speech", "Resist", "Faith",
              "Bravery", "Patience", "Love", "Preach strength", "Preach hit %", "Expected faith / preach",
              "Physical hit % against them", "Status resist %", "Morale surge % / turn", "Ticks per turn"]
    header(ws, 1, labels, [7, 13, 10, 17, 34, 8, 8, 9, 11, 8, 8, 7, 8, 9, 7, 11, 10, 12, 13, 10, 11, 9])
    C = ref
    classes = load_classes_from_csv()
    for row, cls in enumerate(classes.values(), start=2):
        values = [cls["type"], cls["class"], cls["role"], cls["signature"], "|".join(cls["skills"]), cls["speed"], cls["mv"], cls["jump"],
                  cls["mp"], cls["magic_attack"], cls["magic_defense"], cls["faith"], cls["bravery"], cls["patience"], cls["love"]]
        for col, value in enumerate(values, start=1):
            put(ws, row, col, value, INPUT)
        put(ws, row, 16, f"=(L{row}+J{row})*{C['FAITH_TRANSFER_RATE']}", fmt=NUM1)
        put(ws, row, 17, f"=MIN({C['FAITH_MAX_HIT_CHANCE']},MAX({C['FAITH_MIN_HIT_CHANCE']},{C['FAITH_BASE_HIT_CHANCE']}+O{row}*{C['FAITH_HIT_PER_LOVE']}))", fmt=PCT)
        put(ws, row, 18, f"=P{row}*Q{row}", fmt=NUM1)
        put(ws, row, 19, f"=MIN({C['PHYSICAL_MAX_HIT_CHANCE']},MAX({C['PHYSICAL_MIN_HIT_CHANCE']},{C['PHYSICAL_BASE_HIT_CHANCE']}-K{row}*{C['MAGIC_DEFENSE_HIT_SCALE']}))", fmt=PCT)
        put(ws, row, 20, f"=MIN({C['PATIENCE_RESIST_CAP']},N{row}*{C['PATIENCE_RESIST_SCALE']})", fmt=PCT)
        put(ws, row, 21, f"=MIN({C['MORALE_CHANCE_CAP']},{C['MORALE_BASE_CHANCE']}+M{row}*{C['MORALE_CHANCE_SCALE']})", fmt=PCT)
        put(ws, row, 22, f"={C['CT_PER_TURN']}/F{row}", fmt=NUM1)
    last = len(classes) + 1
    # Role averages, to compare the four roles at a glance.
    top = last + 3
    put(ws, top - 1, 1, "Role averages", BOLD)
    header(ws, top, ["Role", "Speed", "Move", "Faith", "Bravery", "Patience", "Love", "Expected faith / preach", "Physical hit % against them", "Status resist %"])
    for i, role in enumerate(("Analyst", "Diplomat", "Sentinel", "Explorer"), start=1):
        r = top + i
        put(ws, r, 1, role, BOLD)
        for col, src in zip(range(2, 11), ["F", "G", "L", "M", "N", "O", "R", "S", "T"]):
            fmt = PCT if src in ("S", "T") else NUM1
            put(ws, r, col, f'=AVERAGEIF($C$2:$C${last},$A{r},{src}$2:{src}${last})', fmt=fmt)
    ws.freeze_panes = "C2"
    return ws, last


# --- Skills ----------------------------------------------------------------------

SKILL_NOTES = {
    "Preach": ("Faith", "Party on enemy: raises their faith by Preach strength; at {FAITH_CAP} they convert. Enemy on party: lowers it. Misses on a Love roll."),
    "Shoot": ("Physical", "Strikes the target down (removed). Archers ignore the uphill penalty. A guard blocks it."),
    "First Aid": ("Heal", "Raises an ally's faith by the user's Preach strength. Always lands."),
    "Heal": ("Heal", "Raises an ally's faith by the user's Preach strength. Always lands."),
    "Lead": ("Support", "+{LEAD_LOVE_BONUS} Love to allies within {LEAD_RADIUS} tiles for {LEAD_TURNS} turns."),
    "Defend": ("Support", "Guards an ally: blocks the next Physical hit."),
    "Slash": ("Physical", "Disarms the target for {DISARM_TURNS} turns (can't act)."),
    "Command": ("Support", "+{COMMAND_CT_BOOST} CT to an ally - their next turn comes sooner."),
    "Fish net": ("Status", "Snares the target for {FISH_NET_SNARE_TURNS} turns (can't move). Patience resists."),
    "Shove": ("Status", "{SHOVE_CHANCE} chance to stun 1 turn and push {SHOVE_DISTANCE} tiles. Patience resists."),
    "Rally": ("Support", "Enemy Medic: lowers an ally's faith by up to {RALLY_FAITH_RESTORE} (not below {RALLY_BASELINE_FAITH}) and clears statuses. Party: raises faith."),
    "Persecute": ("Status", "Never misses: -{PERSECUTE_FAITH} faith and -{PERSECUTE_SLOW} Move for {PERSECUTE_TURNS} turns. At 0 faith the disciple is arrested."),
    "Grand Design": ("Support", "Allies within 2: +1 Move, +5 Speed for 2 turns."),
    "Reason Together": ("Status", "Never misses: +{REASON_TOGETHER_FAITH} faith to an enemy."),
    "Marshal": ("Support", "Allies within 2: +{MARSHAL_CT} CT."),
    "Disputation": ("Status", "Stun 1 turn and +10 faith. Patience resists."),
    "Intercession": ("Support", "Cures an ally's Despair and statuses, +{INTERCESSION_FAITH} faith."),
    "Peacemaker": ("Status", "Disarms an enemy for 2 turns. Patience resists."),
    "Inspire": ("Support", "Allies within 2: +{INSPIRE_FAITH} faith."),
    "Good News": ("Status", "Preaches the target and every enemy beside it at {GOOD_NEWS_SHARE} strength (Love roll each)."),
    "Provision": ("Support", "Refills an ally's MP, cures Snare/Stun, +10 faith."),
    "Shield of Faith": ("Support", "Guards an ally and stops their faith being shaken for 2 turns."),
    "Arrest": ("Status", "Snare {ARREST_SNARE_TURNS} turns and disarm 1. Patience resists."),
    "Breaking Bread": ("Support", "Allies within 1: +{BREAKING_BREAD_FAITH} faith, +10 Love for 2 turns."),
    "Sling": ("Physical", "Ranged disarm (4 tiles) for {DISARM_TURNS} turns."),
    "Psalm": ("Support", "Enemies within 2: -{PSALM_CT} CT. Allies within 2: +10 faith."),
    "Bold Venture": ("Status", "{BOLD_VENTURE_CHANCE} chance: +{BOLD_VENTURE_FAITH} faith to the enemy; otherwise the caster loses {BOLD_VENTURE_BACKFIRE}."),
    "Parable": ("Status", "+{PARABLE_FAITH} faith and -{PARABLE_CT} CT to an enemy, range 3."),
}


def build_skills(wb, ref, units_last, classes_last):
    ws = wb.create_sheet("Skills")
    header(ws, 1, ["Skill", "Type", "MP cost", "Range", "What it does (live numbers)", "Units carrying it", "Classes carrying it"], [17, 10, 9, 8, 95, 13, 13])
    skills = load_skills_from_csv()
    for row, (name, data) in enumerate(skills.items(), start=2):
        put(ws, row, 1, name, BOLD)
        put(ws, row, 2, data["type"], INPUT)
        put(ws, row, 3, data["mp_cost"], INPUT)
        put(ws, row, 4, data["range"], INPUT)
        text = SKILL_NOTES.get(name, (data["type"], ""))[1]
        # Turn {CONSTANT} placeholders into live references to the Constants sheet.
        parts, formula, rest = [], "", text
        while "{" in rest:
            before, _, after = rest.partition("{")
            key, _, rest = after.partition("}")
            parts.append(f'"{before}"')
            value = f"TEXT({ref[key]},\"0%\")" if key in ("SHOVE_CHANCE", "GOOD_NEWS_SHARE", "BOLD_VENTURE_CHANCE") else ref[key]
            parts.append(value)
        parts.append(f'"{rest}"')
        put(ws, row, 5, "=" + "&".join(parts))
        # Exact match on the pipe-separated skill lists, so "Heal" doesn't count "Heal|..." wrongly.
        put(ws, row, 6, f'=SUMPRODUCT(--ISNUMBER(SEARCH("|"&A{row}&"|","|"&Units!$H$2:$H${units_last}&"|")))')
        put(ws, row, 7, f'=SUMPRODUCT(--ISNUMBER(SEARCH("|"&A{row}&"|","|"&Classes!$E$2:$E${classes_last}&"|")))')
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"


# --- Matchup ---------------------------------------------------------------------

def build_matchup(wb, ref, units_last):
    ws = wb.create_sheet("Matchup")
    C = ref
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 70
    ws["A1"] = "Matchup calculator"
    ws["A1"].font = TITLE
    ws["A2"] = "Pick an attacker and a target (yellow cells) - everything below updates."
    ws["A2"].font = Font(name=FONT, italic=True)
    keys = f"Units!$A$2:$A${units_last}"
    dv = DataValidation(type="list", formula1=f"={keys}", allow_blank=False)
    ws.add_data_validation(dv)
    inputs = [("Attacker", "damascus - Paul"), ("Target", "damascus - Ananias"), ("Attacker height above target (tiles)", 0)]
    for i, (label, value) in enumerate(inputs, start=4):
        put(ws, i, 1, label, BOLD)
        cell = put(ws, i, 2, value, INPUT)
        cell.fill = EDIT_FILL
    dv.add("B4")
    dv.add("B5")
    ws["C6"] = "Negative = attacking uphill. Archers ignore the uphill penalty."
    ws["C6"].font = Font(name=FONT, italic=True)

    def look(column, who):
        return f"INDEX(Units!${column}$2:${column}${units_last},MATCH({who},{keys},0))"

    rows = [
        ("Attacker team / class", f'={look("D", "$B$4")}&" / "&{look("E", "$B$4")}', None, ""),
        ("Target team / class", f'={look("D", "$B$5")}&" / "&{look("E", "$B$5")}', None, ""),
        ("Target faith", f"={look('O', '$B$5')}", "0", "Starting faith from the stage's characters.csv."),
        ("PREACH", None, None, None),
        ("Preach strength", f"={look('W', '$B$4')}", NUM1, "Faith moved by one landed Preach."),
        ("Direction", f'=IF(AND({look("D", "$B$4")}<>"Player",{look("D", "$B$5")}="Player"),-1,1)', "0", "+1 raises faith (party preaching enemies, or healing); -1 shakes it (enemies preaching the party)."),
        ("Faith change per landed preach", "=B13*B14", NUM1, ""),
        ("Preach hit chance", f"={look('X', '$B$4')}", PCT, "From the attacker's Love."),
        ("Expected faith change per attempt", "=B15*B16", NUM1, "What a Preach is worth on average."),
        ("Landed preaches to convert (enemy -> cap)", f'=IF(B14>0,ROUNDUP(({C["FAITH_CAP"]}-B11)/B13,0),"-")', "0", "Ignores the 1% faith decay per turn and any Rally."),
        ("Landed preaches to break (faith -> 0)", f'=IF(B14<0,ROUNDUP(B11/B13,0),"-")', "0", "Party member: Despair. A convert: turns back."),
        ("Expected attempts needed", '=IF(ISNUMBER(B18),B18/B16,IF(ISNUMBER(B19),B19/B16,"-"))', NUM1, "Landed preaches / hit chance."),
        ("PHYSICAL (Slash / Shoot / Sling)", None, None, None),
        ("Physical hit chance", f'=MIN({C["PHYSICAL_MAX_HIT_CHANCE"]},MAX({C["PHYSICAL_MIN_HIT_CHANCE"]},{C["PHYSICAL_BASE_HIT_CHANCE"]}'
                                f'+IF({look("E", "$B$4")}="Archer",MAX(0,$B$6),$B$6)*{C["ELEVATION_HIT_BONUS_PER_TILE"]}'
                                f'+IF(OR({look("E", "$B$4")}="Soldier",{look("E", "$B$4")}="Sergeant"),{C["SOLDIER_HIT_BONUS"]},0)'
                                f'-{look("T", "$B$5")}*{C["MAGIC_DEFENSE_HIT_SCALE"]}'
                                f'-IF({look("E", "$B$5")}="Shieldbearer",{C["SHIELDBEARER_HIT_PENALTY"]},0)))', PCT,
         "Includes height, Soldier/Sergeant bonus, target Resist and Shieldbearer penalty. Not the Officer aura or a guard (a guard blocks outright)."),
        ("STATUS", None, None, None),
        ("Target status resist", f"={look('AA', '$B$5')}", PCT, "From the target's Patience."),
        ("Fish net / Arrest / Peacemaker land chance", "=1-B24", PCT, ""),
        ("Shove land chance", f"={C['SHOVE_CHANCE']}*(1-B24)", PCT, "Shove's own chance, then the resist roll."),
        ("Persecutes to arrest", f'=IF({look("D", "$B$5")}="Enemy",ROUNDUP(B11/{C["PERSECUTE_FAITH"]},0),"-")', "0", "Damascus: Persecute never misses."),
    ]
    for i, (label, formula, fmt, note) in enumerate(rows, start=9):
        if formula is None:
            cell = put(ws, i, 1, label, BOLD)
            cell.fill = SECTION_FILL
            ws.cell(row=i, column=2).fill = SECTION_FILL
            ws.cell(row=i, column=3).fill = SECTION_FILL
            continue
        put(ws, i, 1, label)
        put(ws, i, 2, formula, fmt=fmt)
        if note:
            ws.cell(row=i, column=3, value=note).font = Font(name=FONT, italic=True, color="555555")
    return ws


# --- Stage summary ---------------------------------------------------------------

def build_stages(wb, ref, units_last):
    ws = wb.create_sheet("Stage Summary")
    C = ref
    labels = ["Stage", "Title", "Party units", "Enemy units", "Party avg speed", "Enemy avg speed", "Party faith pool", "Enemy faith pool",
              "Party faith / round (expected)", "Enemy faith / round (expected)", "Faith the party must add to convert every enemy",
              "Rounds for the party to convert all", "Rounds for the enemy to break the whole party"]
    header(ws, 1, labels, [11, 28, 8, 8, 9, 9, 10, 10, 13, 13, 16, 13, 14])
    manifest = load_stage_manifest()
    U = lambda col: f"Units!${col}$2:${col}${units_last}"  # noqa: E731
    for row, (node, stage) in enumerate(manifest.items(), start=2):
        put(ws, row, 1, node, INPUT)
        put(ws, row, 2, stage["title"], INPUT)
        put(ws, row, 3, f'=COUNTIFS({U("B")},$A{row},{U("D")},"Player")')
        put(ws, row, 4, f'=COUNTIFS({U("B")},$A{row},{U("D")},"<>Player")')
        put(ws, row, 5, f'=AVERAGEIFS({U("U")},{U("B")},$A{row},{U("D")},"Player")', fmt=NUM1)
        put(ws, row, 6, f'=AVERAGEIFS({U("U")},{U("B")},$A{row},{U("D")},"<>Player")', fmt=NUM1)
        put(ws, row, 7, f'=SUMIFS({U("O")},{U("B")},$A{row},{U("D")},"Player")', fmt="0")
        put(ws, row, 8, f'=SUMIFS({U("O")},{U("B")},$A{row},{U("D")},"<>Player")', fmt="0")
        put(ws, row, 9, f'=SUMIFS({U("Y")},{U("B")},$A{row},{U("D")},"Player")', fmt=NUM1)
        put(ws, row, 10, f'=SUMIFS({U("Y")},{U("B")},$A{row},{U("D")},"<>Player")', fmt=NUM1)
        put(ws, row, 11, f"=D{row}*{C['FAITH_CAP']}-H{row}", fmt="0")
        put(ws, row, 12, f'=IF(I{row}>0,K{row}/I{row},"-")', fmt=NUM1)
        put(ws, row, 13, f'=IF(J{row}>0,G{row}/J{row},"-")', fmt=NUM1)
    last = len(manifest) + 1
    note = last + 2
    ws.cell(row=note, column=1, value="How to read this").font = BOLD
    lines = [
        "A 'round' assumes every unit acts once and spends its turn preaching - a rough measure of faith pressure, not a prediction.",
        "Higher 'Rounds for the party to convert all' = a slower, harder stage for the player; compare it with 'Rounds for the enemy to break the whole party'.",
        "Ignores movement, statuses, Rally, Persecute, physical attacks, escapes and the 1% faith decay. Use tools/playtest_sim.py for full simulations.",
        "Damascus: Saul's band persecutes instead of preaching, so its 'party faith / round' overstates how fast disciples convert there.",
    ]
    for i, text in enumerate(lines, start=1):
        ws.cell(row=note + i, column=1, value=text).font = Font(name=FONT, italic=True)
    ws.freeze_panes = "C2"


# --- Read me ---------------------------------------------------------------------

def build_readme(wb):
    ws = wb.active
    ws.title = "Read Me"
    ws.column_dimensions["A"].width = 120
    lines = [
        ("The Road to Jerusalem - Game Design Balance Sheet", TITLE),
        ("", None),
        ("What this is", BOLD),
        ("Every unit, class, skill and balance constant in the game, with the formulas the game uses written out as live Excel formulas.", None),
        ("Change a number on the Constants sheet (or a unit's stats) and every derived column, matchup and stage summary updates.", None),
        ("This is a sandbox - the game never reads this file. When you like a change, copy it into scripts/game_logic.py or the CSV it came from.", None),
        ("", None),
        ("Sheets", BOLD),
        ("Constants - every tuning number, what it does and where it lives in the code.", None),
        ("Units - every character on every stage: their stats, plus preach strength, hit chances, resist, morale, speed and more.", None),
        ("Classes - the sixteen survey classes, with the same derived numbers and role averages.", None),
        ("Skills - every skill, what it does (with live numbers) and how many units and classes carry it.", None),
        ("Matchup - pick any attacker and target: faith per preach, preaches to convert or break, physical and status odds.", None),
        ("Stage Summary - faith pools and faith pressure per round, side against side, for every stage.", None),
        ("Level Growth - stats gained per level by class; type a level to preview.", None),
        ("", None),
        ("Colours", BOLD),
        ("Blue text - values copied from the game (inputs). Edit these to try ideas.", Font(name=FONT, color="0000FF")),
        ("Black text - formulas. Don't type over them.", None),
        ("Yellow cells - the controls on Matchup and Level Growth.", None),
        ("", None),
        ("Rebuild from the game's current data:  python tools/build_balance_sheet.py   (overwrites this file)", Font(name=FONT, italic=True)),
    ]
    for row, (text, font) in enumerate(lines, start=1):
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = font or Font(name=FONT)
    ws["A20"].fill = EDIT_FILL


def main():
    wb = Workbook()
    build_readme(wb)
    ref = build_constants(wb)
    for name, address in ref.items():
        _CONSTANT_ROWS[name] = int(address.split("$")[-1])
    growth = build_growth(wb)
    units_ws, units_last = build_units(wb, ref, growth)
    classes_ws, classes_last = build_classes(wb, ref)
    build_skills(wb, ref, units_last, classes_last)
    build_matchup(wb, ref, units_last)
    build_stages(wb, ref, units_last)
    # Sheet order: the ones you'll use most first.
    order = ["Read Me", "Matchup", "Units", "Classes", "Skills", "Stage Summary", "Constants", "Level Growth"]
    wb._sheets = [wb[name] for name in order]
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
